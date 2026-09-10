"""Ekspor SEMUA data organisasi → satu workbook Excel (sheet per koleksi) + pembaca balik."""
import io
from datetime import datetime

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from db import db
from data_mgmt_full_schema import (columns_for, org_filter, to_cell, group_of, GROUP_LABELS,
                                   TOO_LARGE)
from data_mgmt_excel import _cell_value

INDEX_SHEET = "_INDEX"
_HEAD = PatternFill("solid", fgColor="1E293B")
_KEY = PatternFill("solid", fgColor="0F766E")
_TYPE = PatternFill("solid", fgColor="F1F5F9")


def sheet_name(coll: str, used: set) -> str:
    base = coll[:31]
    name, i = base, 2
    while name in used:
        suffix = f"~{i}"
        name = base[:31 - len(suffix)] + suffix
        i += 1
    used.add(name)
    return name


def _write_index(ws, rows: list, org: str):
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 34
    ws.column_dimensions["C"].width = 12
    ws.column_dimensions["D"].width = 28
    ws["A1"] = "EKSPOR SEMUA DATA — SIPRO"
    ws["A1"].font = Font(bold=True, size=14, color="0F766E")
    ws["A2"] = f"Organisasi {org} · dibuat {datetime.now():%d %b %Y %H:%M}"
    ws["A3"] = ("Setiap sheet = satu koleksi. Baris 1 = nama field (JANGAN diubah), baris 2 = tipe, "
                "baris 3.. = data. Kolom id = kunci; hapus isi id untuk membuat data baru. "
                f"Sel bertulis {TOO_LARGE} tidak akan diubah saat impor.")
    ws["A3"].alignment = Alignment(wrap_text=True)
    ws.merge_cells("A3:D3")
    ws.row_dimensions[3].height = 48
    for ci, h in enumerate(("sheet", "collection", "count", "group"), start=1):
        c = ws.cell(row=5, column=ci, value=h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = _HEAD
    for ri, r in enumerate(rows, start=6):
        ws.cell(row=ri, column=1, value=r["sheet"])
        ws.cell(row=ri, column=2, value=r["collection"])
        ws.cell(row=ri, column=3, value=r["count"])
        ws.cell(row=ri, column=4, value=GROUP_LABELS[group_of(r["collection"])])


def _write_sheet(ws, cols: list, types: dict, docs: list):
    for ci, k in enumerate(cols, start=1):
        h = ws.cell(row=1, column=ci, value=k)
        h.font = Font(bold=True, color="FFFFFF")
        h.fill = _KEY if k in ("id", "org_id") else _HEAD
        t = ws.cell(row=2, column=ci, value=types[k])
        t.font = Font(italic=True, size=9, color="64748B")
        t.fill = _TYPE
        ws.column_dimensions[get_column_letter(ci)].width = 14 if types[k] in ("int", "float", "bool") else (
            36 if types[k] in ("json", "datetime") or k == "id" else 22)
    ws.freeze_panes = "B3"
    for ri, d in enumerate(docs, start=3):
        for ci, k in enumerate(cols, start=1):
            if k in d:
                ws.cell(row=ri, column=ci, value=to_cell(d[k]))


async def build_full_workbook(org: str, collections: list) -> tuple:
    """→ (bytes, index_rows). collections = daftar nama koleksi yang diekspor."""
    wb = Workbook()
    idx_ws = wb.active
    idx_ws.title = INDEX_SHEET
    used, index_rows = {INDEX_SHEET}, []
    for coll in collections:
        docs = await db[coll].find(org_filter(coll, org)).to_list(None)
        for d in docs:  # _id hanya dipertahankan bila dokumen tak punya id (kunci impor balik)
            if d.get("id"):
                d.pop("_id", None)
            else:
                d["_id"] = str(d["_id"])
        cols, types = columns_for(docs)
        if not cols:
            cols, types = ["id", "org_id"], {"id": "str", "org_id": "str"}
        name = sheet_name(coll, used)
        _write_sheet(wb.create_sheet(name), cols, types, docs)
        index_rows.append({"sheet": name, "collection": coll, "count": len(docs)})
    _write_index(idx_ws, index_rows, org)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue(), index_rows


def parse_full_workbook(content: bytes) -> dict:
    """→ {sheet_name: {collection, columns, types, rows:[{row, cells{col: raw}}]}}"""
    wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    mapping = {}
    if INDEX_SHEET in wb.sheetnames:
        for r in wb[INDEX_SHEET].iter_rows(min_row=6, values_only=True):
            if r and r[0] and r[1]:
                mapping[str(r[0])] = str(r[1])
    out = {}
    for ws in wb.worksheets:
        if ws.title.startswith("_"):  # _INDEX, _TEMUAN (laporan) — bukan data
            continue
        it = ws.iter_rows(values_only=True)
        header = next(it, None)
        if not header:
            continue
        cols = [str(h).strip() if h is not None else "" for h in header]
        type_row = next(it, None) or ()
        types = {}
        known = {"str", "int", "float", "bool", "datetime", "json", "any", "null"}
        has_types = type_row and all((t is None) or str(t).strip() in known for t in type_row)
        if has_types:
            types = {c: str(t).strip() for c, t in zip(cols, type_row) if c and t}
        rows, start = [], 3
        if not has_types and type_row:  # baris 2 sudah data (tipe dihapus pengguna)
            rows.append({"row": 2, "cells": {c: _cell_value(v) for c, v in zip(cols, type_row) if c}})
        for ri, raw in enumerate(it, start=start):
            cells = {c: _cell_value(v) for c, v in zip(cols, raw) if c}
            if any(v is not None for v in cells.values()):
                rows.append({"row": ri, "cells": cells})
        out[ws.title] = {"collection": mapping.get(ws.title, ws.title.strip()),
                         "columns": [c for c in cols if c], "types": types, "rows": rows}
    wb.close()
    return out
