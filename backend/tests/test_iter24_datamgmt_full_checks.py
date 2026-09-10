"""Iter24 — Manajemen Data → Semua Data.
Uji dua fitur baru:
  (1) LAPORAN VALIDASI EXCEL (report.xlsx): sel bermasalah berwarna + komentar,
      sheet _TEMUAN, dan berkas laporan tetap bisa diunggah ulang.
  (2) CEK KONSISTENSI SILANG (checks): Σ termin AR vs unit_total / total /
      harga deal / paid / outstanding — dievaluasi di kondisi SETELAH baris
      sheet diterapkan; sesi tanpa sheet AR/deals/contracts → checks kosong.
"""
import io
import json
import os

import pytest
import requests
from openpyxl import load_workbook

BASE = os.environ["REACT_APP_BACKEND_URL"].rstrip("/")
API = f"{BASE}/api"
PASS = "Sipro#2026"


def _tok(email):
    r = requests.post(f"{API}/auth/login", json={"email": email, "password": PASS}, timeout=20)
    assert r.status_code == 200, r.text
    j = r.json()
    return j.get("access_token") or j.get("token") or (j.get("data") or {}).get("token")


@pytest.fixture(scope="module")
def H():
    return {"Authorization": f"Bearer {_tok('superadmin@sipro.co.id')}"}


def _mongo():
    from pymongo import MongoClient
    from dotenv import dotenv_values
    c = dotenv_values("/app/backend/.env")
    return MongoClient(os.environ.get("MONGO_URL") or c["MONGO_URL"])[
        os.environ.get("DB_NAME") or c["DB_NAME"]]


@pytest.fixture(scope="module")
def ar_context():
    """Pick an AR that has a real deal, prep test targets."""
    sdb = _mongo()
    ars = list(sdb.ar_invoices.find({"org_id": "org-sipro"}, {"_id": 0}))
    assert ars, "Butuh minimal 1 AR di demo"
    # AR yang punya deal
    ar_target = next((a for a in ars if a.get("deal_id") and
                      sdb.deals.find_one({"id": a["deal_id"]})), ars[0])
    # AR lain untuk test AR_VS_DEAL sheet=null
    ar_db_side = next((a for a in ars if a["id"] != ar_target["id"] and a.get("deal_id") and
                       sdb.deals.find_one({"id": a["deal_id"]})), None)
    return {"ars": ars, "target": ar_target, "db_side": ar_db_side}


@pytest.fixture(scope="module")
def session_with_checks(H, ar_context):
    """Build an edited workbook that must trigger multiple cross-checks."""
    r = requests.get(f"{API}/data-mgmt/full/export.xlsx", headers=H,
                     params={"collections": "ar_invoices,deals,customers"}, timeout=60)
    assert r.status_code == 200
    wb = load_workbook(io.BytesIO(r.content))
    ws_ar = wb["ar_invoices"]
    hdr_ar = [c.value for c in ws_ar[1]]
    ci_ar = {h: i + 1 for i, h in enumerate(hdr_ar)}

    target = ar_context["target"]
    db_side = ar_context["db_side"]
    # Pastikan db_side ≠ target dan berbeda dari AR yang akan kita pakai untuk "other"
    ws_d = wb["deals"]
    hdr_d = [c.value for c in ws_d[1]]
    ci_d = {h: i + 1 for i, h in enumerate(hdr_d)}

    # (c) LAKUKAN DULU: buang db_side dari sheet AR + ubah harga dealnya
    if db_side and db_side["id"] != target["id"]:
        for row in range(ws_ar.max_row, 2, -1):
            if ws_ar.cell(row=row, column=ci_ar["id"]).value == db_side["id"]:
                ws_ar.delete_rows(row, 1)
                break
        for row in range(3, ws_d.max_row + 1):
            if ws_d.cell(row=row, column=ci_d["id"]).value == db_side["deal_id"]:
                try:
                    orig_n = float(ws_d.cell(row=row, column=ci_d["price"]).value or 0)
                except Exception:
                    orig_n = 0
                ws_d.cell(row=row, column=ci_d["price"], value=orig_n + 5_000_000)
                break
    else:
        db_side = None

    # Sekarang cari target_row (mungkin bergeser)
    target_row = None
    for row in range(3, ws_ar.max_row + 1):
        if ws_ar.cell(row=row, column=ci_ar["id"]).value == target["id"]:
            target_row = row
            break
    assert target_row, "AR target tidak ditemukan di sheet ekspor"

    # (a) baris target: JSON items → amount termin ke-2 +1_000_000
    items = json.loads(ws_ar.cell(row=target_row, column=ci_ar["items"]).value)
    non_addon_idx = [i for i, it in enumerate(items) if it.get("basis") != "addon"]
    assert len(non_addon_idx) >= 2
    idx = non_addon_idx[1]
    items[idx]["amount"] = float(items[idx].get("amount") or 0) + 1_000_000
    ws_ar.cell(row=target_row, column=ci_ar["items"], value=json.dumps(items))

    # (b) baris AR lain (bukan target, bukan db_side): total=1
    other_row, other_ar_id = None, None
    for row in range(3, ws_ar.max_row + 1):
        rid = ws_ar.cell(row=row, column=ci_ar["id"]).value
        if rid and rid != target["id"] and (not db_side or rid != db_side["id"]):
            other_row = row
            other_ar_id = rid
            break
    if other_row is not None:
        ws_ar.cell(row=other_row, column=ci_ar["total"], value=1)

    buf = io.BytesIO()
    wb.save(buf)
    r = requests.post(f"{API}/data-mgmt/full/sessions", headers=H,
                      files={"file": ("ar_uji_iter24.xlsx", buf.getvalue())}, timeout=60)
    assert r.status_code == 200, r.text
    sid = r.json()["id"]
    yield {"sid": sid, "target_row": target_row, "target_id": target["id"],
           "other_row": other_row, "other_ar_id": other_ar_id,
           "db_side": db_side}
    requests.delete(f"{API}/data-mgmt/full/sessions/{sid}", headers=H, timeout=30)


# ---------- (1) Cross-checks -----------------------------------------------
def test_checks_contain_expected_codes(H, session_with_checks):
    sid = session_with_checks["sid"]
    rep = requests.get(f"{API}/data-mgmt/full/sessions/{sid}", headers=H, timeout=30).json()
    checks = rep.get("checks") or []
    assert isinstance(checks, list) and len(checks) > 0
    # totals.checks == len(checks)
    assert rep["totals"]["checks"] == len(checks)
    # setiap check punya field yang dijanjikan
    for c in checks:
        for k in ("code", "level", "key", "label", "message", "collection"):
            assert k in c, f"missing key {k}: {c}"
        assert c["level"] == "warning"
        assert c["collection"] == "ar_invoices"

    codes_by_key = {}
    for c in checks:
        codes_by_key.setdefault(c["key"], set()).add(c["code"])

    # (a) target row → AR_UNIT_SUM + AR_TOTAL_SUM + AR_VS_DEAL
    target_codes = codes_by_key.get(session_with_checks["target_id"], set())
    assert {"AR_UNIT_SUM", "AR_TOTAL_SUM", "AR_VS_DEAL"} <= target_codes, target_codes
    # dan pointer sheet/row untuk target
    for c in checks:
        if c["key"] == session_with_checks["target_id"]:
            assert c["sheet"] == "ar_invoices"
            assert c["row"] == session_with_checks["target_row"]

    # (b) other row → AR_TOTAL_SUM + AR_OUTSTANDING
    other_codes = codes_by_key.get(session_with_checks["other_ar_id"], set())
    assert {"AR_TOTAL_SUM", "AR_OUTSTANDING"} <= other_codes, other_codes

    # (c) db_side (jika ada) → AR_VS_DEAL di DB, sheet=null, row=null
    db_side = session_with_checks["db_side"]
    if db_side:
        db_checks = [c for c in checks if c["key"] == db_side["id"]]
        assert db_checks and any(c["code"] == "AR_VS_DEAL" for c in db_checks)
        for c in db_checks:
            assert c["sheet"] is None
            assert c["row"] is None


def test_patch_restore_reduces_checks(H, session_with_checks):
    sid = session_with_checks["sid"]
    before = requests.get(f"{API}/data-mgmt/full/sessions/{sid}", headers=H).json()
    before_total = before["totals"]["checks"]
    # Restore items JSON pada target_row menggunakan nilai DB asli.
    sdb = _mongo()
    orig = sdb.ar_invoices.find_one({"id": session_with_checks["target_id"]}, {"_id": 0})
    edits = [{"sheet": "ar_invoices", "row": session_with_checks["target_row"],
              "col": "items", "value": json.dumps(orig["items"])}]
    r = requests.patch(f"{API}/data-mgmt/full/sessions/{sid}/cells", headers=H,
                       json={"edits": edits}, timeout=30)
    assert r.status_code == 200
    after = requests.get(f"{API}/data-mgmt/full/sessions/{sid}", headers=H).json()
    assert after["totals"]["checks"] < before_total, (before_total, after["totals"]["checks"])
    codes = {(c["key"], c["code"]) for c in after["checks"]}
    tid = session_with_checks["target_id"]
    assert (tid, "AR_UNIT_SUM") not in codes
    assert (tid, "AR_TOTAL_SUM") not in codes


# ---------- (2) Report XLSX ------------------------------------------------
def test_report_xlsx_shape_and_reupload(H, session_with_checks):
    sid = session_with_checks["sid"]
    r = requests.get(f"{API}/data-mgmt/full/sessions/{sid}/report.xlsx", headers=H, timeout=60)
    assert r.status_code == 200
    assert "spreadsheet" in r.headers.get("content-type", "").lower() or \
           r.headers.get("content-type", "").startswith(
               "application/vnd.openxmlformats-officedocument.spreadsheetml")
    wb = load_workbook(io.BytesIO(r.content))
    names = wb.sheetnames
    assert names[0] == "_INDEX"
    assert "_TEMUAN" in names
    assert "ar_invoices" in names
    # _TEMUAN header di baris 4
    tw = wb["_TEMUAN"]
    heads = [tw.cell(row=4, column=c).value for c in range(1, 8)]
    assert heads == ["Sheet", "Baris", "Kolom", "Status baris", "Tingkat", "Pesan", "Saran"]
    # Ada minimal 1 temuan cek-silang
    body = [
        [tw.cell(row=r_, column=c).value for c in range(1, 8)]
        for r_ in range(5, tw.max_row + 1)
    ]
    assert body, "harus ada temuan"
    assert any(row[3] == "cek-silang" for row in body), body[:3]
    # Ordering: error → warning → info berdasar sort level
    levels = [row[4] for row in body if row[4]]
    order_map = {"ERROR": 0, "PERINGATAN": 1, "INFO": 2}
    ints = [order_map.get(v, 3) for v in levels]
    assert ints == sorted(ints), levels

    # Sheet data: baris 1 & 2 tetap ada (field & tipe)
    aw = wb["ar_invoices"]
    row1 = [c.value for c in aw[1]]
    row2 = [c.value for c in aw[2]]
    assert "id" in row1 and "org_id" in row1
    assert any(t in ("str", "int", "float", "bool", "datetime", "json") for t in row2 if t)

    # Cari sel error dan sel changed di sheet customers (jika sebelumnya ada
    # error) - di sini fokus di ar_invoices: cari sel dengan fill FECACA atau FDE68A
    from openpyxl.styles import PatternFill  # noqa
    def _fg(cell):
        f = cell.fill.fgColor
        return (f.rgb or "").upper() if f is not None else ""

    err_cells = warn_cells = changed_cells = 0
    for row in aw.iter_rows(min_row=3):
        for cell in row:
            fg = _fg(cell)
            if fg.endswith("FECACA"):
                err_cells += 1
                if cell.comment:
                    assert cell.comment.text  # tidak kosong
            elif fg.endswith("FDE68A"):
                warn_cells += 1
            elif fg.endswith("BBF7D0"):
                changed_cells += 1
    # Kami menyunting sel items dan total → minimal ada sel berubah (BBF7D0)
    assert changed_cells >= 1, "harus ada sel berubah (BBF7D0)"

    # Reupload berkas laporan → sheet _TEMUAN/_INDEX harus dilewati parser
    r2 = requests.post(f"{API}/data-mgmt/full/sessions", headers=H,
                       files={"file": ("laporan.xlsx", r.content)}, timeout=60)
    assert r2.status_code == 200, r2.text
    sid2 = r2.json()["id"]
    try:
        rep2 = requests.get(f"{API}/data-mgmt/full/sessions/{sid2}", headers=H).json()
        sheet_names = {s["sheet"] for s in rep2["sheets"]}
        assert "_TEMUAN" not in sheet_names
        assert "_INDEX" not in sheet_names
        # totals harus sebanding dengan sesi asal (jumlah baris data tetap)
        orig_rep = requests.get(f"{API}/data-mgmt/full/sessions/{sid}", headers=H).json()
        assert rep2["totals"]["rows"] == orig_rep["totals"]["rows"]
    finally:
        requests.delete(f"{API}/data-mgmt/full/sessions/{sid2}", headers=H, timeout=30)


def test_session_without_ar_deal_contract_has_no_checks(H):
    r = requests.get(f"{API}/data-mgmt/full/export.xlsx", headers=H,
                     params={"collections": "customers"}, timeout=60)
    r = requests.post(f"{API}/data-mgmt/full/sessions", headers=H,
                      files={"file": ("c.xlsx", r.content)}, timeout=60)
    sid = r.json()["id"]
    try:
        j = requests.get(f"{API}/data-mgmt/full/sessions/{sid}", headers=H).json()
        assert j["checks"] == []
        assert j["totals"]["checks"] == 0
    finally:
        requests.delete(f"{API}/data-mgmt/full/sessions/{sid}", headers=H, timeout=30)


def test_report_xlsx_missing_session_404(H):
    r = requests.get(f"{API}/data-mgmt/full/sessions/does-not-exist/report.xlsx",
                     headers=H, timeout=30)
    assert r.status_code == 404
