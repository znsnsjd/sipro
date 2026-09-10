"""Uji cepat alur ekspor semua → sunting → sesi impor → validasi → edit sel → commit."""
import io
import os
import sys

import requests
from openpyxl import load_workbook

API = os.environ.get("API", "http://localhost:8001/api")
tok = requests.post(f"{API}/auth/login", json={"email": "superadmin@sipro.co.id", "password": "Sipro#2026"}).json()
token = tok.get("token") or tok.get("access_token") or (tok.get("data") or {}).get("token")
if not token:
    print("login resp", str(tok)[:300])
    sys.exit(1)
H = {"Authorization": f"Bearer {token}"}

r = requests.get(f"{API}/data-mgmt/full/collections", headers=H)
assert r.status_code == 200, r.text
colls = r.json()["collections"]
print("collections", len(colls), "groups", {c["group"] for c in colls})

r = requests.get(f"{API}/data-mgmt/full/export.xlsx", headers=H, params={"collections": "projects,units,customers,counters"})
assert r.status_code == 200, r.text[:300]
wb = load_workbook(io.BytesIO(r.content))
print("sheets", wb.sheetnames)
ws = wb["customers"]
hdr = [c.value for c in ws[1]]
print("customers cols", hdr[:8], "types", [c.value for c in ws[2]][:8])
# sunting: telepon jadi 08xx, org_id salah, id kosong (baru), JSON rusak di units
ci = {h: i + 1 for i, h in enumerate(hdr)}
ws.cell(row=3, column=ci["phone"], value="0812 3456 7890")
ws.cell(row=4, column=ci["org_id"], value="org-lain")
ws.cell(row=5, column=ci["name"], value="Nama Diubah Uji")
last = ws.max_row + 1
ws.cell(row=last, column=ci["name"], value="Pelanggan Baru Uji")
ws.cell(row=last, column=ci["phone"], value="+628111222333")
wu = wb["units"]
uh = {c.value: i + 1 for i, c in enumerate(wu[1])}
jcol = next((k for k, t in zip([c.value for c in wu[1]], [c.value for c in wu[2]]) if t == "json"), None)
print("units json col", jcol)
if jcol:
    wu.cell(row=3, column=uh[jcol], value="{'a': 1}")
wu.cell(row=4, column=uh["price"], value="Rp 1.250.000.000")
wu.cell(row=5, column=uh["status"], value="availble")
buf = io.BytesIO(); wb.save(buf)

r = requests.post(f"{API}/data-mgmt/full/sessions", headers=H, files={"file": ("semua.xlsx", buf.getvalue())})
assert r.status_code == 200, r.text
sid = r.json()["id"]
rep = requests.get(f"{API}/data-mgmt/full/sessions/{sid}", headers=H).json()
print("totals", rep["totals"])
for s in rep["sheets"]:
    print(" ", s["sheet"], s["counts"], "missing", s["missing_count"], s["issues"])

pg = requests.get(f"{API}/data-mgmt/full/sessions/{sid}/sheets/customers", headers=H, params={"only_issues": True}).json()
for row in pg["rows"]:
    print("  cust row", row["row"], row["status"], row["changed"], [i["message"] + (" → " + str(i.get("suggestion")) if "suggestion" in i else "") for i in row["issues"]])
pg = requests.get(f"{API}/data-mgmt/full/sessions/{sid}/sheets/units", headers=H, params={"only_issues": True}).json()
for row in pg["rows"]:
    print("  unit row", row["row"], row["status"], [i["message"] + (" → " + str(i.get("suggestion")) if "suggestion" in i else "") for i in row["issues"]])

# terapkan saran: semua sel yang punya suggestion
edits = []
for sheet in ("customers", "units"):
    pg = requests.get(f"{API}/data-mgmt/full/sessions/{sid}/sheets/{sheet}", headers=H, params={"only_issues": True}).json()
    for row in pg["rows"]:
        for i in row["issues"]:
            if "suggestion" in i and i["level"] in ("error", "warning"):
                edits.append({"sheet": sheet, "row": row["row"], "col": i["col"], "value": i["suggestion"]})
print("edits", len(edits))
r = requests.patch(f"{API}/data-mgmt/full/sessions/{sid}/cells", headers=H, json={"edits": edits})
assert r.status_code == 200, r.text
rep = requests.get(f"{API}/data-mgmt/full/sessions/{sid}", headers=H).json()
print("totals after", rep["totals"])
r = requests.post(f"{API}/data-mgmt/full/sessions/{sid}/commit", headers=H,
                  json={"mode": "update", "sheets": ["customers", "units"], "confirm": "IMPOR"})
print("commit", r.status_code, r.text[:400])
