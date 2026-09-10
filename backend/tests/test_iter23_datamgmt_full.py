"""Iter23 — Manajemen Data → Semua Data (Ekspor/Impor).

Uji end-to-end untuk router /api/data-mgmt/full/*: list koleksi + RBAC, ekspor xlsx,
sesi impor (validasi/saran per sel), edit sel + delete-rows, commit mode update,
commit mode replace pada koleksi kecil (`workers`), dan hapus sesi.
"""
import io
import os
import time

import pytest
import requests
from openpyxl import load_workbook

BASE = os.environ["REACT_APP_BACKEND_URL"].rstrip("/")
API = f"{BASE}/api"
PASS = "Sipro#2026"


def _tok(email):
    r = requests.post(f"{API}/auth/login", json={"email": email, "password": PASS}, timeout=20)
    assert r.status_code == 200, r.text
    j = r.json()
    return j.get("access_token") or j.get("token") or (j.get("data") or {}).get("token")


@pytest.fixture(scope="module")
def H():
    return {"Authorization": f"Bearer {_tok('superadmin@sipro.co.id')}"}


@pytest.fixture(scope="module")
def Hs():
    return {"Authorization": f"Bearer {_tok('sales@sipro.co.id')}"}


# --- 1. Koleksi & RBAC -----------------------------------------------------
def test_collections_ok(H):
    r = requests.get(f"{API}/data-mgmt/full/collections", headers=H, timeout=30)
    assert r.status_code == 200, r.text
    j = r.json()
    assert isinstance(j["collections"], list) and len(j["collections"]) > 20
    c0 = j["collections"][0]
    for k in ("collection", "count", "group", "group_label"):
        assert k in c0
    assert "sessions" in j


def test_collections_forbidden_for_sales(Hs):
    r = requests.get(f"{API}/data-mgmt/full/collections", headers=Hs, timeout=30)
    assert r.status_code == 403


# --- 2. Ekspor xlsx --------------------------------------------------------
def test_export_selected(H):
    r = requests.get(f"{API}/data-mgmt/full/export.xlsx", headers=H,
                     params={"collections": "customers,units"}, timeout=60)
    assert r.status_code == 200
    wb = load_workbook(io.BytesIO(r.content))
    assert "_INDEX" in wb.sheetnames and "customers" in wb.sheetnames and "units" in wb.sheetnames
    hdr = [c.value for c in wb["customers"][1]]
    types = [c.value for c in wb["customers"][2]]
    assert "id" in hdr and "org_id" in hdr
    assert set(types) & {"str", "int", "float", "bool", "datetime", "json"}


def test_export_unknown_collection_400(H):
    r = requests.get(f"{API}/data-mgmt/full/export.xlsx", headers=H,
                     params={"collections": "tidak_ada"}, timeout=30)
    assert r.status_code == 400


# --- 3. Sesi: validasi per sel + saran -------------------------------------
@pytest.fixture(scope="module")
def session_with_issues(H):
    r = requests.get(f"{API}/data-mgmt/full/export.xlsx", headers=H,
                     params={"collections": "customers,units"}, timeout=60)
    wb = load_workbook(io.BytesIO(r.content))
    ws = wb["customers"]
    hdr = [c.value for c in ws[1]]
    ci = {h: i + 1 for i, h in enumerate(hdr)}
    ws.cell(row=3, column=ci["phone"], value="0812 3456 7890")     # info + saran +62
    ws.cell(row=4, column=ci["org_id"], value="org-lain")           # error org
    # baris baru (tanpa id)
    last = ws.max_row + 1
    ws.cell(row=last, column=ci["name"], value="TEST_Pelanggan Baru Iter23")
    ws.cell(row=last, column=ci["phone"], value="+628111000999")

    wu = wb["units"]
    uh = {c.value: i + 1 for i, c in enumerate(wu[1])}
    ut = [c.value for c in wu[2]]
    jcol = next((k for k, t in zip([c.value for c in wu[1]], ut) if t == "json"), None)
    if jcol:
        wu.cell(row=3, column=uh[jcol], value="{'a': 1}")           # error JSON
    wu.cell(row=4, column=uh["price"], value="Rp 1.250.000.000")    # error angka teks
    wu.cell(row=5, column=uh["status"], value="availble")           # warning fuzzy
    buf = io.BytesIO()
    wb.save(buf)
    r = requests.post(f"{API}/data-mgmt/full/sessions", headers=H,
                      files={"file": ("semua.xlsx", buf.getvalue())}, timeout=60)
    assert r.status_code == 200, r.text
    sid = r.json()["id"]
    yield sid
    requests.delete(f"{API}/data-mgmt/full/sessions/{sid}", headers=H, timeout=30)


def _all_issues(H, sid, sheet):
    r = requests.get(f"{API}/data-mgmt/full/sessions/{sid}/sheets/{sheet}",
                     headers=H, params={"only_issues": True, "limit": 500}, timeout=30)
    assert r.status_code == 200
    return r.json()["rows"]


def test_session_report_shape(H, session_with_issues):
    sid = session_with_issues
    r = requests.get(f"{API}/data-mgmt/full/sessions/{sid}", headers=H, timeout=30)
    assert r.status_code == 200
    j = r.json()
    t = j["totals"]
    for k in ("rows", "insert", "update", "same", "error", "warning", "info", "suggestion", "missing"):
        assert k in t
    assert t["error"] >= 3 and t["warning"] >= 1 and t["info"] >= 1
    for s in j["sheets"]:
        assert "rows" not in s  # jangan bocorkan seluruh baris di report


def test_customers_issues_and_suggestions(H, session_with_issues):
    rows = _all_issues(H, session_with_issues, "customers")
    by_row = {r["row"]: r for r in rows}
    # phone info + suggestion +62
    phone = by_row[3]
    iss = [i for i in phone["issues"] if i["col"] == "phone"]
    assert iss and iss[0]["level"] == "info" and iss[0]["suggestion"].startswith("+62")
    # org_id error
    org_row = by_row[4]
    org_iss = [i for i in org_row["issues"] if i["col"] == "org_id"]
    assert org_iss and org_iss[0]["level"] == "error" and org_iss[0]["suggestion"] == "org-sipro"
    # baris baru tanpa id → status insert + info id kosong
    insert_rows = [r for r in rows if r["status"] == "insert"]
    assert insert_rows, "harus ada baris insert (id kosong)"
    assert any(i["level"] == "info" and "id" in i["message"].lower() for i in insert_rows[0]["issues"])


def test_units_issues(H, session_with_issues):
    rows = _all_issues(H, session_with_issues, "units")
    by_row = {r["row"]: r for r in rows}
    # JSON rusak
    j_iss = [i for i in by_row[3]["issues"] if i["level"] == "error" and "JSON" in i["message"]]
    assert j_iss and "suggestion" in j_iss[0]
    # price teks
    p_iss = [i for i in by_row[4]["issues"] if i["col"] == "price"]
    assert p_iss and p_iss[0]["level"] == "error" and p_iss[0]["suggestion"] == 1250000000
    # status fuzzy warning
    s_iss = [i for i in by_row[5]["issues"] if i["col"] == "status"]
    assert s_iss and s_iss[0]["level"] == "warning" and s_iss[0]["suggestion"] == "available"


# --- 4. Terapkan saran → error hilang --------------------------------------
def test_apply_suggestions_removes_errors(H, session_with_issues):
    sid = session_with_issues
    edits = []
    for sheet in ("customers", "units"):
        for row in _all_issues(H, sid, sheet):
            for i in row["issues"]:
                if "suggestion" in i and i["level"] in ("error", "warning"):
                    edits.append({"sheet": sheet, "row": row["row"], "col": i["col"],
                                  "value": i["suggestion"]})
    assert edits, "harus ada saran untuk diterapkan"
    r = requests.patch(f"{API}/data-mgmt/full/sessions/{sid}/cells",
                       headers=H, json={"edits": edits}, timeout=60)
    assert r.status_code == 200
    rep = requests.get(f"{API}/data-mgmt/full/sessions/{sid}", headers=H).json()
    assert rep["totals"]["error"] == 0


# --- 5. Commit confirm salah --------------------------------------------------
def test_commit_bad_confirm_400(H, session_with_issues):
    sid = session_with_issues
    r = requests.post(f"{API}/data-mgmt/full/sessions/{sid}/commit", headers=H,
                      json={"mode": "update", "sheets": ["customers"], "confirm": "IYA"}, timeout=30)
    assert r.status_code == 400


# --- 6. Commit mode update — data benar-benar berubah --------------------------
def test_commit_update_persists(H):
    # ekspor customers saja, ubah nama satu baris, commit, verify via /api/customers
    r = requests.get(f"{API}/data-mgmt/full/export.xlsx", headers=H,
                     params={"collections": "customers"}, timeout=60)
    wb = load_workbook(io.BytesIO(r.content))
    ws = wb["customers"]
    hdr = [c.value for c in ws[1]]
    ci = {h: i + 1 for i, h in enumerate(hdr)}
    cid = ws.cell(row=3, column=ci["id"]).value
    orig_name = ws.cell(row=3, column=ci["name"]).value
    new_name = f"{orig_name} [ITER23]"
    ws.cell(row=3, column=ci["name"], value=new_name)
    buf = io.BytesIO(); wb.save(buf)
    r = requests.post(f"{API}/data-mgmt/full/sessions", headers=H,
                      files={"file": ("c.xlsx", buf.getvalue())}, timeout=60)
    sid = r.json()["id"]
    try:
        r = requests.post(f"{API}/data-mgmt/full/sessions/{sid}/commit", headers=H,
                          json={"mode": "update", "sheets": ["customers"], "confirm": "IMPOR"},
                          timeout=60)
        assert r.status_code == 200, r.text
        res = r.json()
        assert res["totals"]["updated"] >= 1
        assert res["snapshot_before"]["filename"].endswith("__pra-impor-semua-data.json")
        # verifikasi via DB langsung (cust API tidak selalu mengembalikan by id)
        from pymongo import MongoClient
        from dotenv import dotenv_values
        mongo_url = os.environ.get("MONGO_URL") or dotenv_values("/app/backend/.env").get("MONGO_URL")
        db_name = os.environ.get("DB_NAME") or dotenv_values("/app/backend/.env").get("DB_NAME")
        sdb = MongoClient(mongo_url)[db_name]
        got = sdb.customers.find_one({"id": cid, "org_id": "org-sipro"})
        assert got["name"] == new_name
        # kembalikan
        ws.cell(row=3, column=ci["name"], value=orig_name)
        buf2 = io.BytesIO(); wb.save(buf2)
        r = requests.post(f"{API}/data-mgmt/full/sessions", headers=H,
                          files={"file": ("c2.xlsx", buf2.getvalue())}, timeout=60)
        sid2 = r.json()["id"]
        requests.post(f"{API}/data-mgmt/full/sessions/{sid2}/commit", headers=H,
                      json={"mode": "update", "sheets": ["customers"], "confirm": "IMPOR"}, timeout=60)
        requests.delete(f"{API}/data-mgmt/full/sessions/{sid2}", headers=H)
    finally:
        requests.delete(f"{API}/data-mgmt/full/sessions/{sid}", headers=H)


# --- 7. Snapshot muncul di daftar --------------------------------------------
def test_snapshot_listed(H):
    r = requests.get(f"{API}/data-mgmt/snapshots", headers=H, timeout=30)
    assert r.status_code == 200, r.text
    items = r.json() if isinstance(r.json(), list) else r.json().get("data", r.json().get("items", []))
    assert any("pra-impor-semua-data" in (it.get("filename") or "") for it in items)


# --- 8. Commit replace pada koleksi kecil (workers) ---------------------------
def test_commit_replace_workers_deletes_missing(H):
    import os
    from pymongo import MongoClient
    from dotenv import dotenv_values
    mongo_url = os.environ.get("MONGO_URL") or dotenv_values("/app/backend/.env").get("MONGO_URL")
    db_name = os.environ.get("DB_NAME") or dotenv_values("/app/backend/.env").get("DB_NAME")
    sdb = MongoClient(mongo_url)[db_name]
    dummy_id = "TEST_iter23_worker_delete"
    sdb.workers.delete_many({"id": dummy_id})
    sdb.workers.insert_one({
        "id": dummy_id, "org_id": "org-sipro", "name": "TEST_Iter23 Dummy",
        "role": "harian", "daily_wage": 100000})
    assert sdb.workers.count_documents({"id": dummy_id}) == 1
    # ekspor workers
    r = requests.get(f"{API}/data-mgmt/full/export.xlsx", headers=H,
                     params={"collections": "workers"}, timeout=60)
    wb = load_workbook(io.BytesIO(r.content))
    ws = wb["workers"]
    hdr = [c.value for c in ws[1]]
    ci = {h: i + 1 for i, h in enumerate(hdr)}
    # buang baris dummy
    for row in range(ws.max_row, 2, -1):
        if ws.cell(row=row, column=ci["id"]).value == dummy_id:
            ws.delete_rows(row, 1)
            break
    buf = io.BytesIO(); wb.save(buf)
    r = requests.post(f"{API}/data-mgmt/full/sessions", headers=H,
                      files={"file": ("w.xlsx", buf.getvalue())}, timeout=60)
    sid = r.json()["id"]
    try:
        r = requests.post(f"{API}/data-mgmt/full/sessions/{sid}/commit", headers=H,
                          json={"mode": "replace", "sheets": ["workers"], "confirm": "IMPOR"},
                          timeout=60)
        assert r.status_code == 200, r.text
        res = r.json()
        assert res["collections"]["workers"]["deleted"] >= 1
        assert sdb.workers.count_documents({"id": dummy_id}) == 0
    finally:
        requests.delete(f"{API}/data-mgmt/full/sessions/{sid}", headers=H)
        sdb.workers.delete_many({"id": dummy_id})


# --- 9. Delete-rows + delete session -----------------------------------------
def test_delete_rows_and_session(H):
    r = requests.get(f"{API}/data-mgmt/full/export.xlsx", headers=H,
                     params={"collections": "counters"}, timeout=30)
    r = requests.post(f"{API}/data-mgmt/full/sessions", headers=H,
                      files={"file": ("t.xlsx", r.content)}, timeout=30)
    sid = r.json()["id"]
    before = requests.get(f"{API}/data-mgmt/full/sessions/{sid}", headers=H).json()["totals"]["rows"]
    r = requests.post(f"{API}/data-mgmt/full/sessions/{sid}/delete-rows", headers=H,
                      json={"sheet": "counters", "rows": [3]}, timeout=30)
    assert r.status_code == 200
    after = requests.get(f"{API}/data-mgmt/full/sessions/{sid}", headers=H).json()["totals"]["rows"]
    assert after == before - 1
    r = requests.delete(f"{API}/data-mgmt/full/sessions/{sid}", headers=H, timeout=30)
    assert r.status_code == 200
    r = requests.get(f"{API}/data-mgmt/full/sessions/{sid}", headers=H, timeout=30)
    assert r.status_code == 404
