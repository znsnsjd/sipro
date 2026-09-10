"""Laporan validasi sesi impor sebagai Excel: data sheet (tetap bisa diimpor ulang) dengan sel
bermasalah berwarna + komentar pesan/saran, sheet _TEMUAN (daftar semua temuan & cek silang)."""
import io
from datetime import datetime

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from data_mgmt_full_export import INDEX_SHEET, _write_index, _write_sheet
from data_mgmt_full_schema import to_cell

FINDINGS_SHEET = "_TEMUAN"
_FILL = {"error": PatternFill("solid", fgColor="FECACA"), "warning": PatternFill("solid", fgColor="FDE68A"),
         "info": PatternFill("solid", fgColor="BAE6FD"), "changed": PatternFill("solid", fgColor="BBF7D0")}
_FONT = {"error": Font(color="991B1B", bold=True), "warning": Font(color="92400E"),
         "info": Font(color="075985"), "changed": Font(color="065F46")}
_HEAD = PatternFill("solid", fgColor="1E293B")
_LEVEL_LABEL = {"error": "ERROR", "warning": "PERINGATAN", "info": "INFO"}


def _findings_rows(report: dict) -> list:
    rows = []
    for s in report["sheets"]:
        for i in s["issues"]:
            rows.append((s["sheet"], None, None, "sheet", i["level"], i["message"], None))
        for r in s["rows"]:
            for i in r["issues"]:
                rows.append((s["sheet"], r["row"], i["col"], r["status"], i["level"], i["message"],
                             to_cell(i.get("suggestion"))))
    for c in report.get("checks", []):
        rows.append((c.get("sheet") or f"(DB) {c['collection']}", c.get("row"), None, "cek-silang",
                     c["level"], f"{c['label']}: {c['message']}", None))
    order = {"error": 0, "warning": 1, "info": 2}
    rows.sort(key=lambda x: (order.get(x[4], 3), x[0], x[1] or 0))
    return rows


def _write_findings(ws, report: dict):
    ws["A1"] = "TEMUAN VALIDASI"
    ws["A1"].font = Font(bold=True, size=14, color="0F766E")
    t = report["totals"]
    ws["A2"] = (f"{report['filename']} · dibuat {datetime.now():%d %b %Y %H:%M} · {t['rows']} baris · "
                f"{t['error']} error · {t['warning']} peringatan · {t['info']} info · {t['suggestion']} saran · "
                f"{len(report.get('checks', []))} cek silang")
    heads = ("Sheet", "Baris", "Kolom", "Status baris", "Tingkat", "Pesan", "Saran")
    widths = (22, 8, 20, 12, 12, 80, 30)
    for ci, (h, w) in enumerate(zip(heads, widths), start=1):
        c = ws.cell(row=4, column=ci, value=h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = _HEAD
        ws.column_dimensions[get_column_letter(ci)].width = w
    for ri, row in enumerate(_findings_rows(report), start=5):
        for ci, v in enumerate(row, start=1):
            c = ws.cell(row=ri, column=ci, value=v)
            if ci == 5 and v in _FILL:
                c.fill = _FILL[v]
                c.font = _FONT[v]
                c.value = _LEVEL_LABEL[v]
            if ci == 6:
                c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A5"


def _decorate(ws, sheet_rep: dict, cols: list):
    col_idx = {c: i + 1 for i, c in enumerate(cols)}
    for r in sheet_rep["rows"]:
        for c in r["changed"]:
            if c in col_idx:
                cell = ws.cell(row=r["row"], column=col_idx[c])
                cell.fill = _FILL["changed"]
                cell.font = _FONT["changed"]
        by_col = {}
        for i in r["issues"]:
            by_col.setdefault(i["col"], []).append(i)
        for c, issues in by_col.items():
            if c not in col_idx:
                continue
            lvl = "error" if any(i["level"] == "error" for i in issues) else (
                "warning" if any(i["level"] == "warning" for i in issues) else "info")
            cell = ws.cell(row=r["row"], column=col_idx[c])
            cell.fill = _FILL[lvl]
            cell.font = _FONT[lvl]
            lines = [f"[{_LEVEL_LABEL[i['level']]}] {i['message']}"
                     + (f"\n  → saran: {to_cell(i['suggestion'])}" if i.get("suggestion") is not None else "")
                     for i in issues]
            cm = Comment("\n".join(lines)[:2000], "SIPRO")
            cm.width, cm.height = 360, 40 + 40 * len(lines)
            cell.comment = cm


def build_report_workbook(payload: dict, report: dict) -> bytes:
    wb = Workbook()
    idx = wb.active
    idx.title = INDEX_SHEET
    _write_findings(wb.create_sheet(FINDINGS_SHEET), report)
    index_rows = []
    rep_by_sheet = {s["sheet"]: s for s in report["sheets"]}
    for name, s in payload["sheets"].items():
        rep = rep_by_sheet.get(name)
        cols = s["columns"]
        types = (rep or {}).get("types") or {c: s.get("types", {}).get(c, "any") for c in cols}
        ws = wb.create_sheet(name[:31])
        _write_sheet(ws, cols, {c: types.get(c, "any") for c in cols}, [])
        for r in s["rows"]:
            for ci, c in enumerate(cols, start=1):
                v = r["cells"].get(c)
                if v is not None:
                    ws.cell(row=r["row"], column=ci, value=v)
        if rep:
            _decorate(ws, rep, cols)
        index_rows.append({"sheet": name[:31], "collection": s["collection"], "count": len(s["rows"])})
    _write_index(idx, index_rows, payload["org_id"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
